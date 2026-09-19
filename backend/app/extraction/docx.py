import io

from docx import Document

from . import ExtractionError


def extract_docx(data: bytes) -> str:
    try:
        document = Document(io.BytesIO(data))
    except Exception as exc:
        raise ExtractionError(f"No se pudo leer el archivo DOCX: {exc}") from exc

    paragraphs = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.append(cell.text)

    return "\n".join(p for p in paragraphs if p).strip()
